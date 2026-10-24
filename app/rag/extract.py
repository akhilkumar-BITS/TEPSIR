"""Turn an uploaded file into plain text, keeping page numbers where possible."""

import io
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.exceptions import UnsupportedDocumentError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass(slots=True)
class ExtractedPage:
    page_number: int | None
    text: str


def extract_pages(filename: str, content_type: str, data: bytes) -> list[ExtractedPage]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocumentError(
            f"Unsupported file type '{suffix or content_type}'. "
            f"Allowed: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
        )

    if suffix == ".pdf":
        return _extract_pdf(data)
    if suffix == ".docx":
        return _extract_docx(data)
    return _extract_plain_text(data)


def _extract_pdf(data: bytes) -> list[ExtractedPage]:
    reader = PdfReader(io.BytesIO(data))
    pages: list[ExtractedPage] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(ExtractedPage(page_number=index, text=text))
    return pages


def _extract_docx(data: bytes) -> list[ExtractedPage]:
    document = DocxDocument(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    # Tables often carry policy limits or per-diem figures; flatten them too.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    text = "\n\n".join(paragraphs)
    return [ExtractedPage(page_number=None, text=text)] if text else []


def _extract_plain_text(data: bytes) -> list[ExtractedPage]:
    text = data.decode("utf-8", errors="replace").strip()
    return [ExtractedPage(page_number=None, text=text)] if text else []
